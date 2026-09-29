"""Prints the non-empty cells/paragraphs of generated docx files so cell placement can be checked without Word."""
import sys, re, glob, os
import docx
d = sys.argv[1]
for f in sorted(glob.glob(os.path.join(d, '*.docx'))):
    doc = docx.Document(f)
    print('#####', os.path.basename(f))
    for i, p in enumerate(doc.paragraphs):
        t = re.sub(r'[\s　]+', ' ', p.text).strip()
        if t: print(f'  P{i}: {t[:90]}')
    for ti, t in enumerate(doc.tables):
        seen = set()
        for ri, row in enumerate(t.rows):
            for ci, c in enumerate(row.cells):
                if id(c._tc) in seen: continue
                seen.add(id(c._tc))
                txt = re.sub(r'[\s　]+', ' ', c.text).strip()
                if txt: print(f'  T{ti}R{ri}C{ci}: {txt[:90]}')
